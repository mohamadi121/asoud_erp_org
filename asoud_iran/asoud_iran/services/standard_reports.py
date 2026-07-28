from __future__ import annotations

import csv
import hashlib
import io
import json
from decimal import Decimal
from html import escape
from typing import Any

SCHEMA_VERSION = "1.0"
MAX_REPORT_ROWS = 20_000
REPORT_TYPES = {
    "journal",
    "general_ledger",
    "floating_detail_ledger",
    "trial_balance",
    "balance_sheet",
    "profit_and_loss",
}
REPORT_TITLES = {
    "journal": "دفتر روزنامه",
    "general_ledger": "دفتر کل و معین",
    "floating_detail_ledger": "دفتر تفصیلی شناور",
    "trial_balance": "تراز آزمایشی شش‌ستونی",
    "balance_sheet": "صورت وضعیت مالی",
    "profit_and_loss": "صورت سود و زیان",
}


def _number(value: Any) -> float:
    return float(Decimal(str(value or 0)).quantize(Decimal("0.01")))


def _split_balance(value: Any) -> tuple[float, float]:
    amount = _number(value)
    return (amount, 0.0) if amount >= 0 else (0.0, abs(amount))


def canonical_checksum(payload: dict) -> str:
    canonical = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _validate_request(
    company: str,
    report_type: str,
    from_date: str,
    to_date: str,
    branch: str | None,
) -> tuple[str, str, str | None]:
    import frappe
    from frappe.utils import date_diff, getdate

    from asoud_core.permissions import can_access_context

    if report_type not in REPORT_TYPES:
        frappe.throw(f"Unsupported accounting report type: {report_type}")
    start, end = getdate(from_date), getdate(to_date)
    if start > end:
        frappe.throw("Report From Date cannot be after To Date")
    if date_diff(end, start) > 732:
        frappe.throw("A standard accounting report cannot span more than two years")
    branch = branch or None
    if not can_access_context(frappe.session.user, company, branch):
        frappe.throw("Not permitted", frappe.PermissionError)
    if branch and frappe.db.get_value("ASOUD Branch", branch, "company") != company:
        frappe.throw("Report branch does not belong to the selected company")
    return start.isoformat(), end.isoformat(), branch


def _gl_conditions(
    *,
    company: str,
    from_date: str | None = None,
    to_date: str | None = None,
    branch: str | None = None,
    account: str | None = None,
    floating_detail: str | None = None,
) -> tuple[str, dict]:
    conditions = ["gle.company=%(company)s", "gle.is_cancelled=0"]
    values: dict[str, Any] = {"company": company}
    if from_date:
        conditions.append("gle.posting_date >= %(from_date)s")
        values["from_date"] = from_date
    if to_date:
        conditions.append("gle.posting_date <= %(to_date)s")
        values["to_date"] = to_date
    if branch:
        conditions.append("gle.asoud_branch=%(branch)s")
        values["branch"] = branch
    if account:
        conditions.append("gle.account=%(account)s")
        values["account"] = account
    if floating_detail:
        conditions.append("gle.asoud_floating_detail=%(floating_detail)s")
        values["floating_detail"] = floating_detail
    return " AND ".join(conditions), values


def _journal_rows(
    company: str,
    from_date: str,
    to_date: str,
    branch: str | None,
) -> tuple[list[dict], list[dict]]:
    import frappe

    where, values = _gl_conditions(
        company=company,
        from_date=from_date,
        to_date=to_date,
        branch=branch,
    )
    rows = frappe.db.sql(
        f"""
        SELECT
            gle.posting_date,
            ad.temporary_number,
            ad.final_number,
            gle.voucher_type,
            gle.voucher_no,
            acc.account_number,
            gle.account,
            gle.party_type,
            gle.party,
            gle.asoud_detail_code,
            gle.asoud_floating_detail,
            gle.asoud_branch,
            gle.debit,
            gle.credit,
            gle.remarks
        FROM `tabGL Entry` gle
        INNER JOIN `tabAccount` acc ON acc.name=gle.account
        LEFT JOIN `tabASOUD Accounting Document` ad
          ON ad.source_doctype=gle.voucher_type AND ad.source_name=gle.voucher_no
        WHERE {where}
        ORDER BY gle.posting_date, COALESCE(ad.final_number, 2147483647),
                 ad.temporary_number, gle.voucher_no, gle.creation, gle.name
        LIMIT {MAX_REPORT_ROWS + 1}
        """,
        values,
        as_dict=True,
    )
    columns = [
        {"key": "posting_date", "label": "تاریخ", "type": "date"},
        {"key": "temporary_number", "label": "شماره موقت", "type": "text"},
        {"key": "final_number", "label": "شماره قطعی", "type": "integer"},
        {"key": "voucher_type", "label": "نوع سند", "type": "text"},
        {"key": "voucher_no", "label": "شناسه سند", "type": "text"},
        {"key": "account_number", "label": "کد حساب", "type": "text"},
        {"key": "account", "label": "حساب", "type": "text"},
        {"key": "party", "label": "طرف حساب", "type": "text"},
        {"key": "asoud_detail_code", "label": "کد تفصیلی", "type": "text"},
        {"key": "asoud_floating_detail", "label": "تفصیلی شناور", "type": "text"},
        {"key": "asoud_branch", "label": "شعبه", "type": "text"},
        {"key": "debit", "label": "بدهکار", "type": "currency"},
        {"key": "credit", "label": "بستانکار", "type": "currency"},
        {"key": "remarks", "label": "شرح", "type": "text"},
    ]
    return columns, [dict(row) for row in rows]


def _ledger_rows(
    company: str,
    from_date: str,
    to_date: str,
    branch: str | None,
    account: str | None,
    floating_detail: str | None,
    include_detail: bool,
) -> tuple[list[dict], list[dict]]:
    import frappe

    period_where, values = _gl_conditions(
        company=company,
        from_date=from_date,
        to_date=to_date,
        branch=branch,
        account=account,
        floating_detail=floating_detail,
    )
    opening_where, opening_values = _gl_conditions(
        company=company,
        to_date=from_date,
        branch=branch,
        account=account,
        floating_detail=floating_detail,
    )
    opening_where = opening_where.replace(
        "gle.posting_date <= %(to_date)s", "gle.posting_date < %(to_date)s"
    )
    detail_select = (
        "gle.asoud_floating_detail" if include_detail else "NULL AS asoud_floating_detail"
    )
    detail_group = ", gle.asoud_floating_detail" if include_detail else ""
    openings = frappe.db.sql(
        f"""
        SELECT gle.account, {detail_select},
               SUM(gle.debit-gle.credit) balance
          FROM `tabGL Entry` gle
         WHERE {opening_where}
         GROUP BY gle.account{detail_group}
        """,
        opening_values,
        as_dict=True,
    )
    opening_map = {
        (row.account, row.asoud_floating_detail if include_detail else None):
        _number(row.balance)
        for row in openings
    }
    detail_order = (
        "gle.asoud_detail_code, gle.asoud_floating_detail,"
        if include_detail
        else ""
    )
    entries = frappe.db.sql(
        f"""
        SELECT gle.posting_date, ad.temporary_number, ad.final_number,
               gle.voucher_type, gle.voucher_no, acc.account_number,
               gle.account, gle.party, gle.asoud_detail_code,
               gle.asoud_floating_detail, gle.asoud_branch,
               gle.debit, gle.credit, gle.remarks
          FROM `tabGL Entry` gle
          INNER JOIN `tabAccount` acc ON acc.name=gle.account
          LEFT JOIN `tabASOUD Accounting Document` ad
            ON ad.source_doctype=gle.voucher_type AND ad.source_name=gle.voucher_no
         WHERE {period_where}
         ORDER BY acc.account_number, gle.account, {detail_order} gle.posting_date,
                  COALESCE(ad.final_number, 2147483647), gle.creation, gle.name
         LIMIT {MAX_REPORT_ROWS + 1}
        """,
        values,
        as_dict=True,
    )
    running: dict[tuple[str, str | None], float] = {}
    result = []
    for entry in entries:
        key = (
            entry.account,
            entry.asoud_floating_detail if include_detail else None,
        )
        if key not in running:
            running[key] = opening_map.get(key, 0.0)
        running[key] = _number(running[key] + _number(entry.debit) - _number(entry.credit))
        row = dict(entry)
        row["opening_balance"] = opening_map.get(key, 0.0)
        row["running_balance"] = running[key]
        result.append(row)
    columns = [
        {"key": "posting_date", "label": "تاریخ", "type": "date"},
        {"key": "final_number", "label": "شماره قطعی", "type": "integer"},
        {"key": "temporary_number", "label": "شماره موقت", "type": "text"},
        {"key": "voucher_no", "label": "سند", "type": "text"},
        {"key": "account_number", "label": "کد حساب", "type": "text"},
        {"key": "account", "label": "حساب", "type": "text"},
    ]
    if include_detail:
        columns.extend(
            [
                {"key": "asoud_detail_code", "label": "کد تفصیلی", "type": "text"},
                {
                    "key": "asoud_floating_detail",
                    "label": "تفصیلی شناور",
                    "type": "text",
                },
            ]
        )
    columns.extend(
        [
            {"key": "opening_balance", "label": "مانده اول دوره", "type": "currency"},
            {"key": "debit", "label": "بدهکار", "type": "currency"},
            {"key": "credit", "label": "بستانکار", "type": "currency"},
            {"key": "running_balance", "label": "مانده", "type": "currency"},
            {"key": "asoud_branch", "label": "شعبه", "type": "text"},
            {"key": "remarks", "label": "شرح", "type": "text"},
        ]
    )
    return columns, result


def _trial_balance_rows(
    company: str,
    from_date: str,
    to_date: str,
    branch: str | None,
) -> tuple[list[dict], list[dict]]:
    import frappe

    branch_condition = "AND gle.asoud_branch=%(branch)s" if branch else ""
    values = {
        "company": company,
        "from_date": from_date,
        "to_date": to_date,
        "branch": branch,
    }
    balances = frappe.db.sql(
        f"""
        SELECT acc.account_number, gle.account, acc.root_type,
               SUM(CASE WHEN gle.posting_date < %(from_date)s
                        THEN gle.debit-gle.credit ELSE 0 END) opening_balance,
               SUM(CASE WHEN gle.posting_date BETWEEN %(from_date)s AND %(to_date)s
                        THEN gle.debit ELSE 0 END) period_debit,
               SUM(CASE WHEN gle.posting_date BETWEEN %(from_date)s AND %(to_date)s
                        THEN gle.credit ELSE 0 END) period_credit
          FROM `tabGL Entry` gle
          INNER JOIN `tabAccount` acc ON acc.name=gle.account
         WHERE gle.company=%(company)s AND gle.posting_date<=%(to_date)s
           AND gle.is_cancelled=0 {branch_condition}
         GROUP BY acc.account_number, gle.account, acc.root_type
         ORDER BY acc.account_number, gle.account
        """,
        values,
        as_dict=True,
    )
    rows = []
    for balance in balances:
        opening_debit, opening_credit = _split_balance(balance.opening_balance)
        closing = (
            _number(balance.opening_balance)
            + _number(balance.period_debit)
            - _number(balance.period_credit)
        )
        closing_debit, closing_credit = _split_balance(closing)
        rows.append(
            {
                "account_number": balance.account_number,
                "account": balance.account,
                "root_type": balance.root_type,
                "opening_debit": opening_debit,
                "opening_credit": opening_credit,
                "period_debit": _number(balance.period_debit),
                "period_credit": _number(balance.period_credit),
                "closing_debit": closing_debit,
                "closing_credit": closing_credit,
            }
        )
    columns = [
        {"key": "account_number", "label": "کد حساب", "type": "text"},
        {"key": "account", "label": "حساب", "type": "text"},
        {"key": "opening_debit", "label": "اول دوره بدهکار", "type": "currency"},
        {"key": "opening_credit", "label": "اول دوره بستانکار", "type": "currency"},
        {"key": "period_debit", "label": "گردش بدهکار", "type": "currency"},
        {"key": "period_credit", "label": "گردش بستانکار", "type": "currency"},
        {"key": "closing_debit", "label": "مانده بدهکار", "type": "currency"},
        {"key": "closing_credit", "label": "مانده بستانکار", "type": "currency"},
    ]
    return columns, rows


def _financial_statement_rows(
    company: str,
    report_type: str,
    from_date: str,
    to_date: str,
    branch: str | None,
) -> tuple[list[dict], list[dict]]:
    import frappe

    if report_type == "balance_sheet":
        roots = ("Asset", "Liability", "Equity")
        date_condition = "gle.posting_date<=%(to_date)s"
    else:
        roots = ("Income", "Expense")
        date_condition = "gle.posting_date BETWEEN %(from_date)s AND %(to_date)s"
    branch_condition = "AND gle.asoud_branch=%(branch)s" if branch else ""
    rows = frappe.db.sql(
        f"""
        SELECT acc.root_type, acc.account_number, gle.account,
               SUM(gle.debit-gle.credit) raw_balance
          FROM `tabGL Entry` gle
          INNER JOIN `tabAccount` acc ON acc.name=gle.account
         WHERE gle.company=%(company)s AND gle.is_cancelled=0
           AND {date_condition} AND acc.root_type IN %(roots)s {branch_condition}
         GROUP BY acc.root_type, acc.account_number, gle.account
         ORDER BY FIELD(acc.root_type, 'Asset', 'Liability', 'Equity', 'Income', 'Expense'),
                  acc.account_number, gle.account
        """,
        {
            "company": company,
            "from_date": from_date,
            "to_date": to_date,
            "roots": roots,
            "branch": branch,
        },
        as_dict=True,
    )
    result = []
    for row in rows:
        raw = _number(row.raw_balance)
        amount = -raw if row.root_type in {"Liability", "Equity", "Income"} else raw
        result.append(
            {
                "section": row.root_type,
                "account_number": row.account_number,
                "account": row.account,
                "amount": _number(amount),
            }
        )
    columns = [
        {"key": "section", "label": "گروه", "type": "text"},
        {"key": "account_number", "label": "کد حساب", "type": "text"},
        {"key": "account", "label": "حساب", "type": "text"},
        {"key": "amount", "label": "مبلغ", "type": "currency"},
    ]
    return columns, result


def _totals(report_type: str, rows: list[dict]) -> dict:
    if report_type in {"journal", "general_ledger", "floating_detail_ledger"}:
        debit = _number(sum(_number(row.get("debit")) for row in rows))
        credit = _number(sum(_number(row.get("credit")) for row in rows))
        return {
            "debit": debit,
            "credit": credit,
            "difference": _number(debit - credit),
        }
    if report_type == "trial_balance":
        keys = (
            "opening_debit",
            "opening_credit",
            "period_debit",
            "period_credit",
            "closing_debit",
            "closing_credit",
        )
        return {key: _number(sum(_number(row.get(key)) for row in rows)) for key in keys}
    section_totals: dict[str, float] = {}
    for row in rows:
        section_totals[row["section"]] = _number(
            section_totals.get(row["section"], 0) + _number(row["amount"])
        )
    if report_type == "balance_sheet":
        section_totals["difference"] = _number(
            section_totals.get("Asset", 0)
            - section_totals.get("Liability", 0)
            - section_totals.get("Equity", 0)
        )
    else:
        section_totals["net_profit"] = _number(
            section_totals.get("Income", 0) - section_totals.get("Expense", 0)
        )
    return section_totals


def build_standard_report(
    *,
    company: str,
    report_type: str,
    from_date: str,
    to_date: str,
    branch: str | None = None,
    account: str | None = None,
    floating_detail: str | None = None,
) -> dict:
    import frappe

    from asoud_iran.accounting.jalali import format_jalali

    from_date, to_date, branch = _validate_request(
        company, report_type, from_date, to_date, branch
    )
    if account:
        account_company = frappe.db.get_value("Account", account, "company")
        if account_company != company:
            frappe.throw("Report account does not belong to the selected company")
    if floating_detail and not frappe.db.exists(
        "ASOUD Floating Detail Company",
        {"parent": floating_detail, "company": company, "enabled": 1},
    ):
        frappe.throw("Floating detail is not enabled for the selected company")

    if report_type == "journal":
        columns, rows = _journal_rows(company, from_date, to_date, branch)
    elif report_type in {"general_ledger", "floating_detail_ledger"}:
        columns, rows = _ledger_rows(
            company,
            from_date,
            to_date,
            branch,
            account,
            floating_detail if report_type == "floating_detail_ledger" else None,
            report_type == "floating_detail_ledger",
        )
    elif report_type == "trial_balance":
        columns, rows = _trial_balance_rows(company, from_date, to_date, branch)
    else:
        columns, rows = _financial_statement_rows(
            company, report_type, from_date, to_date, branch
        )
    if len(rows) > MAX_REPORT_ROWS:
        frappe.throw(
            f"Report exceeds {MAX_REPORT_ROWS} rows; narrow the date, branch or account filters"
        )
    totals = _totals(report_type, rows)
    company_meta = frappe.db.get_value(
        "Company",
        company,
        ["company_name", "tax_id", "asoud_national_id", "asoud_economic_code"],
        as_dict=True,
    )
    payload = {
        "schema_version": SCHEMA_VERSION,
        "report_type": report_type,
        "title": REPORT_TITLES[report_type],
        "company": company,
        "company_identity": dict(company_meta or {}),
        "branch": branch,
        "from_date": from_date,
        "to_date": to_date,
        "from_date_jalali": format_jalali(from_date),
        "to_date_jalali": format_jalali(to_date),
        "currency": "IRR",
        "filters": {
            "account": account,
            "floating_detail": floating_detail,
        },
        "columns": columns,
        "rows": rows,
        "totals": totals,
        "row_count": len(rows),
    }
    payload["checksum"] = canonical_checksum(payload)
    return payload


def render_csv(report: dict) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([report["title"]])
    writer.writerow(["شرکت", report["company"]])
    writer.writerow(["شعبه", report.get("branch") or "همه"])
    writer.writerow(["از تاریخ", report["from_date_jalali"]])
    writer.writerow(["تا تاریخ", report["to_date_jalali"]])
    writer.writerow(["واحد", report["currency"]])
    writer.writerow([])
    columns = report["columns"]
    writer.writerow([column["label"] for column in columns])
    for row in report["rows"]:
        writer.writerow([row.get(column["key"], "") for column in columns])
    writer.writerow([])
    writer.writerow(["جمع/کنترل"])
    for key, value in report["totals"].items():
        writer.writerow([key, value])
    writer.writerow(["checksum", report["checksum"]])
    return b"\xef\xbb\xbf" + output.getvalue().encode("utf-8")


def render_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ASOUD"
    sheet.sheet_view.rightToLeft = True
    sheet.append([report["title"]])
    sheet.append(["شرکت", report["company"]])
    sheet.append(["شعبه", report.get("branch") or "همه"])
    sheet.append(["از تاریخ", report["from_date_jalali"]])
    sheet.append(["تا تاریخ", report["to_date_jalali"]])
    sheet.append(["واحد", report["currency"]])
    sheet.append([])
    columns = report["columns"]
    sheet.append([column["label"] for column in columns])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2457C5")
        cell.alignment = Alignment(horizontal="center")
    for row in report["rows"]:
        sheet.append([row.get(column["key"], "") for column in columns])
    sheet.append([])
    sheet.append(["جمع/کنترل"])
    for key, value in report["totals"].items():
        sheet.append([key, value])
    sheet.append(["checksum", report["checksum"]])
    for column_cells in sheet.columns:
        length = min(
            45,
            max((len(str(cell.value or "")) for cell in column_cells), default=10) + 2,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = length
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_pdf(report: dict) -> bytes:
    from frappe.utils.pdf import get_pdf

    columns = report["columns"]
    header = "".join(f"<th>{escape(column['label'])}</th>" for column in columns)
    body = "".join(
        "<tr>"
        + "".join(
            f"<td>{escape(str(row.get(column['key'], '') or ''))}</td>"
            for column in columns
        )
        + "</tr>"
        for row in report["rows"]
    )
    totals = "".join(
        f"<tr><th>{escape(str(key))}</th><td>{escape(str(value))}</td></tr>"
        for key, value in report["totals"].items()
    )
    html = f"""<!doctype html>
<html lang="fa" dir="rtl"><head><meta charset="utf-8">
<style>
@page {{ size: A4 landscape; margin: 12mm; }}
body {{ font-family: Tahoma, Arial, sans-serif; direction: rtl; font-size: 8pt; }}
h1,p {{ text-align: right; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 8px; }}
th,td {{ border: 1px solid #9aa9b8; padding: 4px; text-align: right; }}
th {{ background: #eaf2f8; }}
.meta {{ width: 60%; }}
</style></head><body>
<h1>{escape(report["title"])}</h1>
<p>شرکت: {escape(report["company"])} | شعبه: {escape(report.get("branch") or "همه")}
 | دوره: {escape(report["from_date_jalali"])} تا {escape(report["to_date_jalali"])}
 | واحد: {escape(report["currency"])}</p>
<table><thead><tr>{header}</tr></thead><tbody>{body}</tbody></table>
<table class="meta">{totals}<tr><th>Checksum</th><td>{report["checksum"]}</td></tr></table>
</body></html>"""
    return get_pdf(html)


def render_export(report: dict, file_format: str) -> tuple[bytes, str, str]:
    normalized = file_format.lower()
    stem = (
        f"asoud-{report['report_type']}-{report['company']}-"
        f"{report['from_date']}-{report['to_date']}"
    ).replace(" ", "-")
    if normalized == "csv":
        return render_csv(report), f"{stem}.csv", "text/csv; charset=utf-8"
    if normalized in {"xlsx", "excel"}:
        return (
            render_xlsx(report),
            f"{stem}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    if normalized == "pdf":
        return render_pdf(report), f"{stem}.pdf", "application/pdf"
    raise ValueError("Export format must be csv, xlsx or pdf")
