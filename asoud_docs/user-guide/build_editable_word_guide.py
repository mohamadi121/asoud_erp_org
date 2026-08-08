from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from lxml import html
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "ASOUD_ERP_SETUP_USER_GUIDE_FA.html"
OUTPUT = ROOT / "ASOUD_ERP_SETUP_USER_GUIDE_FA_EDITABLE.docx"

FONT_NAME = "Tahoma"
TEXT_COLOR = "17233A"
HEADING_COLOR = "102A56"
MUTED_COLOR = "53627A"
BLUE = "1769E0"
LIGHT_BLUE = "F3F7FD"
LIGHT_BORDER = "DCE3EE"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_BORDER, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def set_table_rtl(table) -> None:
    properties = table._tbl.tblPr
    bidi = properties.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        properties.append(bidi)


def set_paragraph_rtl(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = alignment
    properties = paragraph._p.get_or_add_pPr()
    bidi = properties.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        properties.append(bidi)
    bidi.set(qn("w:val"), "1")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.18


def set_run_font(run, size: float = 10, bold: bool = False, color: str = TEXT_COLOR) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._r.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)
    rtl = run_properties.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        run_properties.append(rtl)


def add_text(paragraph, text: str, *, size: float = 10, bold: bool = False, color: str = TEXT_COLOR):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def normalized_text(element) -> str:
    return " ".join(" ".join(element.itertext()).split())


def add_inline_content(paragraph, element, *, size: float = 10, color: str = TEXT_COLOR) -> None:
    if element.text and element.text.strip():
        add_text(paragraph, element.text.strip(), size=size, color=color)
    for child in element:
        text = normalized_text(child)
        if text:
            add_text(
                paragraph,
                text,
                size=size,
                bold=child.tag in {"b", "strong"},
                color=BLUE if child.tag in {"code", "kbd"} else color,
            )
        if child.tail and child.tail.strip():
            add_text(paragraph, child.tail.strip(), size=size, color=color)


def add_heading(container, text: str, level: int) -> None:
    paragraph = container.add_paragraph()
    paragraph.style = f"Heading {min(level, 3)}"
    size = {1: 22, 2: 16, 3: 12}.get(level, 11)
    add_text(paragraph, text, size=size, bold=True, color=HEADING_COLOR)
    set_paragraph_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    paragraph.paragraph_format.space_after = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.keep_with_next = True


def add_paragraph_from_element(container, element, *, size: float = 10, color: str = TEXT_COLOR):
    paragraph = container.add_paragraph()
    add_inline_content(paragraph, element, size=size, color=color)
    set_paragraph_rtl(paragraph)
    return paragraph


def add_list(container, element, ordered: bool) -> None:
    for index, item in enumerate(element.xpath("./li"), start=1):
        paragraph = container.add_paragraph()
        prefix = f"{index}. " if ordered else "• "
        add_text(paragraph, prefix, size=10, bold=True, color=BLUE)
        add_inline_content(paragraph, item, size=10)
        set_paragraph_rtl(paragraph)
        paragraph.paragraph_format.right_indent = Cm(0.4)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)


def image_width(path: Path, maximum: float = 4.85) -> Inches:
    with Image.open(path) as image:
        width, height = image.size
    if width <= 0 or height <= 0:
        return Inches(maximum)
    return Inches(maximum)


def add_image(container, element, maximum: float = 4.85) -> None:
    source = element.get("src", "")
    path = (SOURCE.parent / source).resolve()
    if not path.exists():
        paragraph = container.add_paragraph()
        add_text(paragraph, f"[تصویر یافت نشد: {source}]", color="B42318")
        set_paragraph_rtl(paragraph)
        return
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=image_width(path, maximum))
    paragraph.paragraph_format.space_after = Pt(4)


def add_html_table(container, element) -> None:
    rows = element.xpath("./tr|./thead/tr|./tbody/tr")
    if not rows:
        return
    column_count = max(len(row.xpath("./th|./td")) for row in rows)
    table = container.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_rtl(table)
    for row_index, source_row in enumerate(rows):
        cells = table.add_row().cells
        source_cells = source_row.xpath("./th|./td")
        for column_index, source_cell in enumerate(source_cells):
            cell = cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index == 0 or source_cell.tag == "th":
                set_cell_shading(cell, "EDF3FB")
            paragraph = cell.paragraphs[0]
            add_inline_content(
                paragraph,
                source_cell,
                size=9,
                color=HEADING_COLOR if row_index == 0 or source_cell.tag == "th" else TEXT_COLOR,
            )
            for run in paragraph.runs:
                if row_index == 0 or source_cell.tag == "th":
                    run.font.bold = True
            set_paragraph_rtl(paragraph)
    container.add_paragraph()


def add_card(container, element, fill: str = LIGHT_BLUE) -> None:
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_rtl(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    inline_tags = {"b", "strong", "span", "code", "kbd"}
    paragraph = None
    if element.text and element.text.strip():
        paragraph = cell.add_paragraph()
        add_text(paragraph, element.text.strip(), size=9.5, color=MUTED_COLOR)
    for child in element:
        if child.tag in inline_tags:
            if paragraph is None:
                paragraph = cell.add_paragraph()
            add_text(
                paragraph,
                normalized_text(child),
                size=9.5,
                bold=child.tag in {"b", "strong"},
                color=BLUE if child.tag in {"code", "kbd"} else MUTED_COLOR,
            )
            if child.tail and child.tail.strip():
                add_text(paragraph, child.tail.strip(), size=9.5, color=MUTED_COLOR)
            continue
        if paragraph is not None:
            set_paragraph_rtl(paragraph)
            paragraph = None
        render_element(cell, child)
        if child.tail and child.tail.strip():
            tail_paragraph = cell.add_paragraph()
            add_text(tail_paragraph, child.tail.strip(), size=9.5, color=MUTED_COLOR)
            set_paragraph_rtl(tail_paragraph)
    if paragraph is not None:
        set_paragraph_rtl(paragraph)
    if len(cell.paragraphs) == 0:
        empty = cell.add_paragraph()
        set_paragraph_rtl(empty)
    container.add_paragraph()


def add_two_column(container, element) -> None:
    children = [child for child in element if isinstance(child.tag, str)]
    if not children:
        return
    table = container.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_rtl(table)
    cells = table.rows[0].cells
    for cell in cells:
        cell.width = Inches(5.05)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_border(cell, color="FFFFFF", size="0")
        for paragraph in cell.paragraphs:
            paragraph._element.getparent().remove(paragraph._element)
    for index, child in enumerate(children):
        target = cells[min(index, 1)]
        render_element(target, child)
    container.add_paragraph()


def render_element(container, element) -> None:
    if not isinstance(element.tag, str):
        return
    tag = element.tag.lower()
    classes = set((element.get("class") or "").split())
    if tag in {"h1", "h2", "h3"}:
        add_heading(container, normalized_text(element), int(tag[1]))
    elif tag == "p":
        add_paragraph_from_element(
            container,
            element,
            size=10.5 if "lead" in classes else 9 if "caption" in classes else 10,
            color=MUTED_COLOR if classes & {"lead", "caption"} else TEXT_COLOR,
        )
    elif tag == "img":
        add_image(container, element)
    elif tag == "table":
        add_html_table(container, element)
    elif tag == "ul":
        add_list(container, element, ordered=False)
    elif tag == "ol":
        add_list(container, element, ordered=True)
    elif tag == "div" and "grid-2" in classes:
        add_two_column(container, element)
    elif tag == "div" and ("grid-equal" in classes or "mini-images" in classes):
        add_two_column(container, element)
    elif tag == "div" and ("page-intro" in classes or "card" in classes):
        fill = "EEF5FF"
        if "ok" in classes:
            fill = "EDF9F3"
        elif "warn" in classes:
            fill = "FFF8E8"
        elif "future" in classes:
            fill = "F5F0FF"
        add_card(container, element, fill=fill)
    elif tag == "div" and "parts" in classes:
        paragraph = container.add_paragraph()
        labels = [normalized_text(child) for child in element if normalized_text(child)]
        add_text(paragraph, "  |  ".join(labels), size=8.5, bold=True, color="31547C")
        set_paragraph_rtl(paragraph)
    elif tag == "div" and "step" in classes:
        paragraph = container.add_paragraph()
        pieces = [normalized_text(child) for child in element if normalized_text(child)]
        if pieces:
            add_text(paragraph, f"{pieces[0]}  ", size=10, bold=True, color=BLUE)
            add_text(paragraph, " ".join(pieces[1:]), size=10)
        set_paragraph_rtl(paragraph)
        paragraph.paragraph_format.right_indent = Cm(0.25)
    elif tag == "div" and "checklist" in classes:
        for child in element:
            paragraph = container.add_paragraph()
            add_text(paragraph, "✓ ", size=10, bold=True, color="15935D")
            add_text(paragraph, normalized_text(child), size=10)
            set_paragraph_rtl(paragraph)
    elif tag == "div" and "toc" in classes:
        for child in element:
            paragraph = container.add_paragraph()
            add_text(paragraph, normalized_text(child), size=10)
            set_paragraph_rtl(paragraph)
    elif tag in {"div", "section"}:
        if element.text and element.text.strip() and not list(element):
            add_paragraph_from_element(container, element)
        else:
            for child in element:
                render_element(container, child)


def add_page_number(paragraph) -> None:
    set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    add_text(paragraph, "آسود ERP  •  راهنمای صفحات نهایی  •  صفحه ", size=8, color="7B879A")
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])
    set_run_font(run, size=8, color="7B879A")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.45)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    style_properties = normal.element.get_or_add_rPr()
    fonts = style_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        style_properties.append(fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        heading_style = document.styles[style_name]
        heading_style.font.name = FONT_NAME
        heading_style.font.color.rgb = RGBColor.from_string(HEADING_COLOR)
        heading_properties = heading_style.element.get_or_add_rPr()
        heading_fonts = heading_properties.find(qn("w:rFonts"))
        if heading_fonts is None:
            heading_fonts = OxmlElement("w:rFonts")
            heading_properties.append(heading_fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            heading_fonts.set(qn(f"w:{attribute}"), FONT_NAME)

    add_page_number(section.footer.paragraphs[0])


def render_cover(document: Document, section_element) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_rtl(table)
    cell = table.cell(0, 0)
    cell.width = Inches(10.2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "123567")
    set_cell_border(cell, color="123567", size="0")
    for paragraph in cell.paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    title = section_element.xpath(".//h1")
    subtitle = section_element.xpath(".//h2")
    paragraphs = section_element.xpath("./div/p")
    for _ in range(3):
        cell.add_paragraph()
    if title:
        paragraph = cell.add_paragraph()
        add_text(paragraph, normalized_text(title[0]), size=28, bold=True, color="FFFFFF")
        set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    if subtitle:
        paragraph = cell.add_paragraph()
        add_text(paragraph, normalized_text(subtitle[0]), size=16, color="D9E8FF")
        set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    for source in paragraphs:
        paragraph = cell.add_paragraph()
        add_text(paragraph, normalized_text(source), size=10, color="FFFFFF")
        set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        cell.add_paragraph()


def build() -> None:
    tree = html.fromstring(SOURCE.read_text(encoding="utf-8"))
    pages = tree.xpath("//section[contains(concat(' ', normalize-space(@class), ' '), ' page ')]")
    document = Document()
    configure_document(document)

    for page_index, page in enumerate(pages):
        if page_index:
            document.add_page_break()
        if "cover" in set((page.get("class") or "").split()):
            render_cover(document, page)
            continue
        for child in page:
            render_element(document, child)

    properties = document.core_properties
    properties.title = "راهنمای کامل صفحات نهایی آسود ERP"
    properties.subject = "راهنمای کاربری قابل ویرایش"
    properties.author = "ASOUD ERP"
    properties.keywords = "ASOUD, ERP, راهنمای کاربری, تنظیمات, کارتابل, حسابداری"
    document.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Pages/sections: {len(pages)}")


if __name__ == "__main__":
    build()
